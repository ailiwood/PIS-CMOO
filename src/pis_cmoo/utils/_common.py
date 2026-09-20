"""
_common.py — shared utilities for G3.5/G4 scripts (real-data analysis).

Centralised paths, hashing, and loaders so every downstream script reads the
same frozen artifacts and computes SHA256 identically. Read-only with respect
to the canonical source xlsx and the frozen CSV.
"""
from __future__ import annotations

import hashlib
import json
from datetime import datetime, timezone
from pathlib import Path

import numpy as np
import pandas as pd
import yaml

PROJECT_ROOT = Path(__file__).resolve().parents[2]
WORKSPACE = PROJECT_ROOT / "010_true_data_analysis"
CONFIG_DIR = WORKSPACE / "00_config"
AUDIT_DIR = WORKSPACE / "01_audit"
DATA_DIR = WORKSPACE / "02_data"
FROZEN_DIR = DATA_DIR / "frozen"
MANIFEST_DIR = DATA_DIR / "manifests"
INTERIM_DIR = DATA_DIR / "interim"
SPLITS_DIR = WORKSPACE / "03_splits"
PRED_DIR = WORKSPACE / "04_prediction"
UQ_DIR = WORKSPACE / "05_uq"
INTERP_DIR = WORKSPACE / "06_interpretation"
OPT_DIR = WORKSPACE / "07_optimization"
VALIDATION_DIR = WORKSPACE / "08_validation"
FIGURES_DIR = WORKSPACE / "09_figures"
TABLES_DIR = WORKSPACE / "10_tables"
MANUSCRIPT_DIR = WORKSPACE / "11_manuscript_support"
HANDOFFS_DIR = WORKSPACE / "12_handoffs"
SCRIPTS_DIR = WORKSPACE / "scripts"
TESTS_DIR = WORKSPACE / "tests"
LOGS_DIR = WORKSPACE / "logs"

CANONICAL_SOURCE = PROJECT_ROOT / "data" / "true_data08221157.xlsx"
FROZEN_CSV = FROZEN_DIR / "real_data_analysis_v1.csv"

EXPECTED_SOURCE_SHA256 = "31e6236cf8fc4792c07fc11c9ae78f1eb0b89dd34baa5f855707cc2026912492"
EXPECTED_FROZEN_SHA256 = "f46e6794eeb1357788d892528fe1bd2ba26df54d45c17fe8ddb7625ab41481cc"

SEED_MASTER = 20260822

# 29 columns in the canonical source (26 raw + 3 derived)
SOURCE_COLUMNS = [
    "province", "mine_id", "project_id", "borehole_id", "record_id",
    "mine_type", "lithology", "karst_class", "fracture_class",
    "hydrogeological_regime", "project_stage",
    "depth_m", "section_length_m", "pre_lugeon", "groundwater_pressure_MPa",
    "WCR", "GP_MPa", "grout_take_L_per_m", "grouting_time_h",
    "allowable_GP_MPa", "min_WCR", "max_WCR", "max_grout_take_L_per_m",
    "max_grouting_time_h", "post_lugeon", "acceptance_threshold_lu",
    "delta_lugeon", "relative_reduction", "log_reduction",
]

# 26 raw (non-derived) columns — must be exactly preserved under NaN-equivalence
RAW_COLUMNS = SOURCE_COLUMNS[:-3]
DERIVED_COLUMNS = ["delta_lugeon", "relative_reduction", "log_reduction"]

CORE_REQUIRED = ["depth_m", "section_length_m", "pre_lugeon", "WCR", "GP_MPa",
                 "grout_take_L_per_m", "grouting_time_h", "post_lugeon",
                 "acceptance_threshold_lu"]

EXTENDED_CONTEXT_COLUMNS = ["groundwater_pressure_MPa", "lithology", "karst_class",
                            "fracture_class", "hydrogeological_regime",
                            "mine_type", "project_stage"]

DEVELOPMENT_MINES = ["GX-M01", "GX-M02", "GZ-M01", "GZ-M02", "YN-M01", "YN-M02"]
EXTERNAL_MINE = "SC-M01"
ALL_MINES = sorted(DEVELOPMENT_MINES + [EXTERNAL_MINE])

DECISION_COLS = ["WCR", "GP_MPa", "grout_take_L_per_m", "grouting_time_h"]
NUMERIC_CONTEXT_CORE = ["depth_m", "section_length_m", "pre_lugeon", "groundwater_pressure_MPa"]
CATEGORICAL_CONTEXT = ["lithology", "karst_class", "fracture_class", "hydrogeological_regime",
                       "mine_type", "project_stage"]
ENGINEERING_CONSTRAINT_COLS = ["allowable_GP_MPa", "min_WCR", "max_WCR",
                               "max_grout_take_L_per_m", "max_grouting_time_h"]


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def now_utc() -> str:
    return datetime.now(timezone.utc).isoformat()


def read_frozen() -> pd.DataFrame:
    """Read the frozen analysis CSV (source of truth for all downstream work)."""
    return pd.read_csv(FROZEN_CSV)


def read_source_raw() -> pd.DataFrame:
    """Read the canonical xlsx via openpyxl, normalise BOM/header, drop empty rows."""
    from openpyxl import load_workbook
    wb = load_workbook(CANONICAL_SOURCE, read_only=True, data_only=True)
    ws = wb["Sheet1"]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(x).lstrip("﻿").strip() if x is not None else None for x in rows[0]]
    df = pd.DataFrame(rows[1:], columns=header)
    if df.columns[0] != "province":
        df = df.rename(columns={df.columns[0]: "province"})
    df = df[df["record_id"].notna()].copy()
    return df


def load_yaml(name: str) -> dict:
    with (CONFIG_DIR / name).open("r", encoding="utf-8") as f:
        return yaml.safe_load(f)


def write_json(path: Path, payload: dict) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    return path


def s_number(record_id: str) -> int:
    """Extract the integer S-section index from a record_id like '...-S03'."""
    return int(str(record_id).rsplit("-S", 1)[-1])


def is_missing_scalar(x) -> bool:
    """Treat None, NaN, and empty/whitespace string as missing."""
    if x is None:
        return True
    if isinstance(x, float) and np.isnan(x):
        return True
    if isinstance(x, str) and x.strip() == "":
        return True
    return False


def canonical_value(x):
    """Normalise a cell to ('MISS',), ('NUM', float_rounded), or ('STR', stripped)."""
    if is_missing_scalar(x):
        return ("MISS",)
    if isinstance(x, str):
        try:
            return ("NUM", round(float(x), 12))
        except ValueError:
            return ("STR", x.strip())
    try:
        return ("NUM", round(float(x), 12))
    except (TypeError, ValueError):
        return ("STR", str(x).strip())


def ensure_dirs():
    for d in [AUDIT_DIR, MANIFEST_DIR, INTERIM_DIR, SPLITS_DIR, PRED_DIR, UQ_DIR,
              INTERP_DIR, OPT_DIR, VALIDATION_DIR, FIGURES_DIR, TABLES_DIR,
              MANUSCRIPT_DIR, HANDOFFS_DIR, SCRIPTS_DIR, TESTS_DIR, LOGS_DIR]:
        d.mkdir(parents=True, exist_ok=True)
